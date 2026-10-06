import csv
import io
import json
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape
import re


CSV_HEADERS = [
    "rank",
    "insight_type",
    "title",
    "source_columns",
    "method",
    "verification_status",
    "confidence_level",
    "confidence_score",
    "priority_score",
    "evidence",
    "limitations",
    "explanation_status",
    "explanation",
    "recommendation_action",
    "recommendation_reason",
    "recommendation_priority",
]


def _insight_key(insight: dict) -> tuple:
    return (
        insight.get("insight_type"),
        insight.get("method"),
        tuple(insight.get("source_columns", [])),
        insight.get("title"),
    )


def _spreadsheet_safe(value: Any) -> Any:
    """
    Prevent CSV / spreadsheet formula injection.

    Numeric values remain numeric. Strings beginning with spreadsheet
    formula-control characters are prefixed with an apostrophe.
    """
    if value is None:
        return ""

    if isinstance(value, (int, float, bool)):
        return value

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )
    else:
        value = str(value)

    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value

    return value


def _build_lookup(items: list[dict] | None) -> dict:
    return {
        _insight_key(item): item
        for item in (items or [])
    }


def build_verified_insight_rows(
    ranked_insights: list[dict] | None,
    explained_insights: list[dict] | None = None,
    recommended_insights: list[dict] | None = None,
) -> list[dict]:
    """
    Build export rows strictly from already-ranked verified insights.

    Raw analysis results and uploaded dataset rows are intentionally
    not accepted by this function.
    """
    explanations = _build_lookup(
        explained_insights
    )

    recommendations = _build_lookup(
        recommended_insights
    )

    rows = []

    for rank, insight in enumerate(
        ranked_insights or [],
        start=1,
    ):
        key = _insight_key(insight)

        explained = explanations.get(
            key,
            {},
        )

        recommended = recommendations.get(
            key,
            {},
        )

        explanation = explained.get(
            "explanation",
            {},
        ) or {}

        recommendation = recommended.get(
            "recommendation",
            {},
        ) or {}

        recommendation_content = (
            recommendation.get(
                "recommendation",
                {},
            )
            if isinstance(
                recommendation.get(
                    "recommendation"
                ),
                dict,
            )
            else recommendation
        )

        confidence = insight.get(
            "confidence",
            {},
        ) or {}

        verification = insight.get(
            "verification",
            {},
        ) or {}

        rows.append({
            "rank": rank,
            "insight_type": insight.get(
                "insight_type"
            ),
            "title": insight.get(
                "title"
            ),
            "source_columns": ", ".join(
                str(column)
                for column in insight.get(
                    "source_columns",
                    [],
                )
            ),
            "method": insight.get(
                "method"
            ),
            "verification_status": (
                verification.get(
                    "status"
                )
            ),
            "confidence_level": (
                confidence.get(
                    "level"
                )
            ),
            "confidence_score": (
                confidence.get(
                    "score"
                )
            ),
            "priority_score": (
                insight.get(
                    "score"
                )
            ),
            "evidence": insight.get(
                "evidence",
                {},
            ),
            "limitations": insight.get(
                "limitations",
                [],
            ),
            "explanation_status": (
                explanation.get(
                    "status"
                )
                if isinstance(
                    explanation,
                    dict,
                )
                else None
            ),
            "explanation": (
                explanation.get(
                    "explanation",
                    {},
                ).get(
                    "text"
                )
                if isinstance(
                    explanation,
                    dict,
                )
                and isinstance(
                    explanation.get(
                        "explanation"
                    ),
                    dict,
                )
                else ""
            ),
            "recommendation_action": (
                recommendation_content.get(
                    "action"
                )
                if isinstance(
                    recommendation_content,
                    dict,
                )
                else ""
            ),
            "recommendation_reason": (
                recommendation_content.get(
                    "reason"
                )
                if isinstance(
                    recommendation_content,
                    dict,
                )
                else ""
            ),
            "recommendation_priority": (
                recommendation_content.get(
                    "priority"
                )
                if isinstance(
                    recommendation_content,
                    dict,
                )
                else ""
            ),
        })

    return rows


def build_csv_export(
    ranked_insights: list[dict] | None,
    explained_insights: list[dict] | None = None,
    recommended_insights: list[dict] | None = None,
) -> bytes:
    rows = build_verified_insight_rows(
        ranked_insights,
        explained_insights,
        recommended_insights,
    )

    buffer = io.StringIO(
        newline=""
    )

    writer = csv.DictWriter(
        buffer,
        fieldnames=CSV_HEADERS,
    )

    writer.writeheader()

    for row in rows:
        writer.writerow({
            key: _spreadsheet_safe(
                row.get(key)
            )
            for key in CSV_HEADERS
        })

    return buffer.getvalue().encode(
        "utf-8-sig"
    )


def _write_table(
    worksheet,
    headers,
    rows,
):
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True
        )

    for row in rows:
        worksheet.append([
            _spreadsheet_safe(
                value
            )
            for value in row
        ])

    worksheet.freeze_panes = "A2"

    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    for column_index, column in enumerate(
        worksheet.columns,
        start=1,
    ):
        width = 10

        for cell in column:
            value = cell.value

            if value is None:
                continue

            width = max(
                width,
                min(
                    len(str(value)) + 2,
                    48,
                ),
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = width


def build_xlsx_export(
    *,
    run_id,
    dataset_version_id,
    status,
    created_at,
    ranked_insights,
    explained_insights,
    recommended_insights,
) -> bytes:
    rows = build_verified_insight_rows(
        ranked_insights,
        explained_insights,
        recommended_insights,
    )

    workbook = Workbook()

    summary = workbook.active
    summary.title = "Executive Summary"

    high = sum(
        1
        for row in rows
        if row["confidence_level"]
        == "high"
    )

    medium = sum(
        1
        for row in rows
        if row["confidence_level"]
        == "medium"
    )

    low = sum(
        1
        for row in rows
        if row["confidence_level"]
        == "low"
    )

    top = rows[0] if rows else {}

    summary_rows = [
        (
            "Verified signals",
            len(rows),
        ),
        (
            "High confidence",
            high,
        ),
        (
            "Medium confidence",
            medium,
        ),
        (
            "Low confidence",
            low,
        ),
        (
            "Top finding",
            top.get(
                "title",
                "No ranked insight",
            ),
        ),
        (
            "Top priority score",
            top.get(
                "priority_score",
                "",
            ),
        ),
    ]

    _write_table(
        summary,
        [
            "Metric",
            "Value",
        ],
        summary_rows,
    )

    verified = workbook.create_sheet(
        "Verified Insights"
    )

    verified_headers = [
        "Rank",
        "Insight Type",
        "Title",
        "Source Columns",
        "Method",
        "Verification",
        "Confidence Level",
        "Confidence Score",
        "Priority Score",
        "Limitations",
        "Grounded Explanation",
        "Recommendation",
        "Recommendation Reason",
        "Recommendation Priority",
    ]

    verified_rows = [
        [
            row["rank"],
            row["insight_type"],
            row["title"],
            row["source_columns"],
            row["method"],
            row["verification_status"],
            row["confidence_level"],
            row["confidence_score"],
            row["priority_score"],
            row["limitations"],
            row["explanation"],
            row["recommendation_action"],
            row["recommendation_reason"],
            row["recommendation_priority"],
        ]
        for row in rows
    ]

    _write_table(
        verified,
        verified_headers,
        verified_rows,
    )

    evidence_sheet = (
        workbook.create_sheet(
            "Evidence"
        )
    )

    evidence_rows = []

    for row in rows:
        evidence = row.get(
            "evidence",
            {},
        )

        if not isinstance(
            evidence,
            dict,
        ):
            continue

        for field, value in evidence.items():
            evidence_rows.append([
                row["rank"],
                row["title"],
                field,
                value,
            ])

    _write_table(
        evidence_sheet,
        [
            "Rank",
            "Insight",
            "Evidence Field",
            "Evidence Value",
        ],
        evidence_rows,
    )

    metadata = workbook.create_sheet(
        "Run Metadata"
    )

    metadata_rows = [
        (
            "Analysis Run ID",
            str(run_id),
        ),
        (
            "Dataset Version ID",
            str(dataset_version_id),
        ),
        (
            "Run Status",
            status,
        ),
        (
            "Created At",
            (
                created_at.isoformat()
                if hasattr(
                    created_at,
                    "isoformat",
                )
                else created_at
            ),
        ),
        (
            "Export Scope",
            (
                "Ranked verified insights, "
                "grounded explanations, "
                "recommendations, and run metadata only"
            ),
        ),
        (
            "Raw Dataset Included",
            "No",
        ),
        (
            "Protected / PII Columns Included",
            "No",
        ),
    ]

    _write_table(
        metadata,
        [
            "Field",
            "Value",
        ],
        metadata_rows,
    )

    output = io.BytesIO()
    workbook.save(output)

    return output.getvalue()


def _plain_report_text(
    value: Any,
) -> str:
    if value is None:
        return ""

    if isinstance(
        value,
        (dict, list, tuple),
    ):
        value = json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )

    value = str(value)

    value = re.sub(
        r"(?m)^#{1,6}\s*",
        "",
        value,
    )

    value = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        value,
    )

    value = re.sub(
        r"(?m)^\s*[-*]\s+",
        "- ",
        value,
    )

    value = value.replace(
        "---",
        "",
    )

    return value.strip()


def _pdf_paragraph(
    value: Any,
    style,
):
    text = _plain_report_text(
        value
    )

    safe = escape(text)

    safe = safe.replace(
        "\n",
        "<br/>",
    )

    return Paragraph(
        safe or "-",
        style,
    )


def build_pdf_export(
    *,
    run_id,
    dataset_version_id,
    status,
    created_at,
    ranked_insights,
    explained_insights,
    recommended_insights,
) -> bytes:
    """
    Build a presentation-ready PDF exclusively
    from persisted verified reporting outputs.

    Raw dataset rows and protected columns are
    intentionally outside this function boundary.
    """
    rows = build_verified_insight_rows(
        ranked_insights,
        explained_insights,
        recommended_insights,
    )

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=17 * mm,
        leftMargin=17 * mm,
        topMargin=19 * mm,
        bottomMargin=18 * mm,
        title=(
            "Signal Ledger "
            "Verified Analysis Report"
        ),
        author="Signal Ledger",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=27,
        textColor=colors.HexColor(
            "#16312C"
        ),
        spaceAfter=6,
        alignment=TA_LEFT,
    )

    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor(
            "#667B74"
        ),
        spaceAfter=15,
    )

    heading_style = ParagraphStyle(
        "ReportHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=colors.HexColor(
            "#16312C"
        ),
        spaceBefore=8,
        spaceAfter=8,
    )

    insight_title_style = ParagraphStyle(
        "InsightTitle",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor(
            "#1F3933"
        ),
        spaceAfter=5,
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor(
            "#344B45"
        ),
        spaceAfter=6,
    )

    label_style = ParagraphStyle(
        "ReportLabel",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor(
            "#5A7069"
        ),
    )

    small_style = ParagraphStyle(
        "ReportSmall",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=7.2,
        leading=10,
        textColor=colors.HexColor(
            "#667B74"
        ),
    )

    story = []

    story.append(
        Paragraph(
            "Signal Ledger",
            title_style,
        )
    )

    story.append(
        Paragraph(
            (
                "Verified Business "
                "Analytics Report"
            ),
            subtitle_style,
        )
    )

    created_text = (
        created_at.isoformat()
        if hasattr(
            created_at,
            "isoformat",
        )
        else str(
            created_at or ""
        )
    )

    metadata_data = [
        [
            _pdf_paragraph(
                "Analysis Run",
                label_style,
            ),
            _pdf_paragraph(
                str(run_id),
                small_style,
            ),
        ],
        [
            _pdf_paragraph(
                "Dataset Version",
                label_style,
            ),
            _pdf_paragraph(
                str(
                    dataset_version_id
                ),
                small_style,
            ),
        ],
        [
            _pdf_paragraph(
                "Status",
                label_style,
            ),
            _pdf_paragraph(
                status,
                small_style,
            ),
        ],
        [
            _pdf_paragraph(
                "Created",
                label_style,
            ),
            _pdf_paragraph(
                created_text,
                small_style,
            ),
        ],
    ]

    metadata_table = Table(
        metadata_data,
        colWidths=[
            34 * mm,
            135 * mm,
        ],
    )

    metadata_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (0, -1),
                colors.HexColor(
                    "#EFF4F1"
                ),
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor(
                    "#D8E1DD"
                ),
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP",
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6,
            ),
        ])
    )

    story.append(
        metadata_table
    )

    story.append(
        Spacer(
            1,
            7 * mm,
        )
    )

    story.append(
        Paragraph(
            "Executive Summary",
            heading_style,
        )
    )

    high = sum(
        row[
            "confidence_level"
        ] == "high"
        for row in rows
    )

    medium = sum(
        row[
            "confidence_level"
        ] == "medium"
        for row in rows
    )

    low = sum(
        row[
            "confidence_level"
        ] == "low"
        for row in rows
    )

    summary_data = [
        [
            _pdf_paragraph(
                "Verified signals",
                label_style,
            ),
            _pdf_paragraph(
                len(rows),
                body_style,
            ),
            _pdf_paragraph(
                "High confidence",
                label_style,
            ),
            _pdf_paragraph(
                high,
                body_style,
            ),
        ],
        [
            _pdf_paragraph(
                "Medium confidence",
                label_style,
            ),
            _pdf_paragraph(
                medium,
                body_style,
            ),
            _pdf_paragraph(
                "Low confidence",
                label_style,
            ),
            _pdf_paragraph(
                low,
                body_style,
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            39 * mm,
            20 * mm,
            39 * mm,
            20 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, -1),
                colors.HexColor(
                    "#F7F9F8"
                ),
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor(
                    "#D9E2DE"
                ),
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP",
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
        ])
    )

    story.append(
        summary_table
    )

    if rows:
        story.append(
            Spacer(
                1,
                4 * mm,
            )
        )

        story.append(
            Paragraph(
                "Top Verified Finding",
                heading_style,
            )
        )

        top = rows[0]

        top_table = Table(
            [
                [
                    _pdf_paragraph(
                        top[
                            "title"
                        ],
                        insight_title_style,
                    )
                ],
                [
                    _pdf_paragraph(
                        (
                            f"Confidence: "
                            f"{top['confidence_level'] or '-'}"
                            f" | Priority score: "
                            f"{top['priority_score']}"
                        ),
                        body_style,
                    )
                ],
            ],
            colWidths=[
                169 * mm
            ],
        )

        top_table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor(
                        "#EEF5E8"
                    ),
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    0.6,
                    colors.HexColor(
                        "#B9CCAC"
                    ),
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    10,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    10,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
            ])
        )

        story.append(
            top_table
        )

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    story.append(
        Paragraph(
            "Verified Findings",
            heading_style,
        )
    )

    if not rows:
        story.append(
            Paragraph(
                (
                    "No ranked verified "
                    "insights were available "
                    "for this run."
                ),
                body_style,
            )
        )

    for row in rows:
        content = []

        content.append(
            Paragraph(
                (
                    f"{row['rank']}. "
                    f"{escape(str(row['title']))}"
                ),
                insight_title_style,
            )
        )

        content.append(
            _pdf_paragraph(
                (
                    f"Type: "
                    f"{row['insight_type']} | "
                    f"Confidence: "
                    f"{row['confidence_level']} | "
                    f"Confidence score: "
                    f"{row['confidence_score']} | "
                    f"Priority score: "
                    f"{row['priority_score']}"
                ),
                small_style,
            )
        )

        content.append(
            _pdf_paragraph(
                (
                    "Source columns: "
                    f"{row['source_columns']}"
                ),
                small_style,
            )
        )

        evidence = row.get(
            "evidence",
            {},
        )

        if isinstance(
            evidence,
            dict,
        ) and evidence:
            evidence_data = [
                [
                    _pdf_paragraph(
                        "Evidence",
                        label_style,
                    ),
                    _pdf_paragraph(
                        "Verified value",
                        label_style,
                    ),
                ]
            ]

            for key, value in (
                evidence.items()
            ):
                evidence_data.append([
                    _pdf_paragraph(
                        key.replace(
                            "_",
                            " ",
                        ),
                        small_style,
                    ),
                    _pdf_paragraph(
                        value,
                        small_style,
                    ),
                ])

            evidence_table = Table(
                evidence_data,
                colWidths=[
                    58 * mm,
                    111 * mm,
                ],
                repeatRows=1,
            )

            evidence_table.setStyle(
                TableStyle([
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#EFF4F1"
                        ),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.3,
                        colors.HexColor(
                            "#D9E1DE"
                        ),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ])
            )

            content.append(
                Spacer(
                    1,
                    2 * mm,
                )
            )

            content.append(
                evidence_table
            )

        if row.get(
            "explanation"
        ):
            content.append(
                Spacer(
                    1,
                    2.5 * mm,
                )
            )

            content.append(
                Paragraph(
                    "Grounded Explanation",
                    label_style,
                )
            )

            content.append(
                _pdf_paragraph(
                    row[
                        "explanation"
                    ],
                    body_style,
                )
            )

        if row.get(
            "recommendation_action"
        ):
            content.append(
                Paragraph(
                    "Recommended Next Step",
                    label_style,
                )
            )

            content.append(
                _pdf_paragraph(
                    row[
                        "recommendation_action"
                    ],
                    body_style,
                )
            )

        if row.get(
            "recommendation_reason"
        ):
            content.append(
                _pdf_paragraph(
                    row[
                        "recommendation_reason"
                    ],
                    small_style,
                )
            )

        limitations = row.get(
            "limitations",
            [],
        )

        if limitations:
            content.append(
                Paragraph(
                    "Limitations",
                    label_style,
                )
            )

            if isinstance(
                limitations,
                list,
            ):
                limitations_text = (
                    "\n".join(
                        f"- {item}"
                        for item
                        in limitations
                    )
                )
            else:
                limitations_text = (
                    str(
                        limitations
                    )
                )

            content.append(
                _pdf_paragraph(
                    limitations_text,
                    small_style,
                )
            )

        story.append(
            KeepTogether(
                content
            )
        )

        story.append(
            Spacer(
                1,
                5 * mm,
            )
        )

    story.append(
        Spacer(
            1,
            3 * mm,
        )
    )

    trust_table = Table(
        [[
            _pdf_paragraph(
                (
                    "Trust boundary: this report "
                    "contains persisted ranked "
                    "verified insights, grounded "
                    "explanations, recommendations, "
                    "and run metadata only. "
                    "Raw dataset rows and protected "
                    "PII columns are not included."
                ),
                small_style,
            )
        ]],
        colWidths=[
            169 * mm
        ],
    )

    trust_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, -1),
                colors.HexColor(
                    "#F4F7F5"
                ),
            ),
            (
                "BOX",
                (0, 0),
                (-1, -1),
                0.5,
                colors.HexColor(
                    "#D5DEDA"
                ),
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                8,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                8,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
        ])
    )

    story.append(
        trust_table
    )

    def footer(
        canvas,
        doc,
    ):
        canvas.saveState()

        canvas.setStrokeColor(
            colors.HexColor(
                "#DCE4E0"
            )
        )

        canvas.line(
            17 * mm,
            12 * mm,
            A4[0] - 17 * mm,
            12 * mm,
        )

        canvas.setFont(
            "Helvetica",
            7,
        )

        canvas.setFillColor(
            colors.HexColor(
                "#788C85"
            )
        )

        canvas.drawString(
            17 * mm,
            7.5 * mm,
            "Signal Ledger - Verified analytics",
        )

        canvas.drawRightString(
            A4[0] - 17 * mm,
            7.5 * mm,
            f"Page {doc.page}",
        )

        canvas.restoreState()

    document.build(
        story,
        onFirstPage=footer,
        onLaterPages=footer,
    )

    return output.getvalue()
