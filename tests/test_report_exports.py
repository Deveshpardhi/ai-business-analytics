import csv
import io
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook

from app.services.report_exporter import (
    build_csv_export,
    build_pdf_export,
    build_verified_insight_rows,
    build_xlsx_export,
)


def make_insight(
    *,
    title="Revenue shows a rising trend",
):
    return {
        "insight_type": (
            "time_series_trend"
        ),
        "title": title,
        "source_columns": [
            "Revenue",
            "Date",
        ],
        "method": (
            "deterministic_time_series_trend"
        ),
        "evidence": {
            "percentage_change": 25.0,
            "sample_size": 12,
        },
        "verification": {
            "status": "verified",
        },
        "confidence": {
            "level": "high",
            "score": 0.91,
        },
        "score": 0.88,
        "limitations": [
            (
                "Historical trends do not "
                "forecast future results."
            )
        ],
    }


def make_explained(
    insight,
):
    return {
        **insight,
        "explanation": {
            "status": "approved",
            "explanation": {
                "text": (
                    "Revenue increased "
                    "over the observed period."
                )
            },
        },
    }


def make_recommended(
    insight,
):
    return {
        **insight,
        "recommendation": {
            "status": "generated",
            "recommendation": {
                "action": (
                    "Monitor the trend."
                ),
                "reason": (
                    "Historical movement "
                    "requires investigation."
                ),
                "priority": "medium",
            },
        },
    }


def test_verified_export_rows_merge_trusted_outputs():
    insight = make_insight()

    rows = (
        build_verified_insight_rows(
            [insight],
            [
                make_explained(
                    insight
                )
            ],
            [
                make_recommended(
                    insight
                )
            ],
        )
    )

    assert len(rows) == 1

    row = rows[0]

    assert row["rank"] == 1

    assert (
        row["verification_status"]
        == "verified"
    )

    assert (
        row["confidence_score"]
        == 0.91
    )

    assert (
        row["explanation"]
        == (
            "Revenue increased "
            "over the observed period."
        )
    )

    assert (
        row[
            "recommendation_action"
        ]
        == "Monitor the trend."
    )


def test_csv_export_contains_ranked_insights_only():
    insight = make_insight()

    content = build_csv_export(
        [insight],
        [
            make_explained(
                insight
            )
        ],
        [
            make_recommended(
                insight
            )
        ],
    )

    decoded = content.decode(
        "utf-8-sig"
    )

    reader = csv.DictReader(
        io.StringIO(decoded)
    )

    rows = list(reader)

    assert len(rows) == 1

    assert (
        rows[0]["title"]
        == insight["title"]
    )

    assert (
        "raw_customer_email"
        not in decoded
    )


def test_csv_export_prevents_formula_injection():
    insight = make_insight(
        title="=HYPERLINK(\"bad\")"
    )

    content = build_csv_export(
        [insight]
    )

    decoded = content.decode(
        "utf-8-sig"
    )

    rows = list(
        csv.DictReader(
            io.StringIO(
                decoded
            )
        )
    )

    assert rows[0]["title"].startswith(
        "'="
    )


def test_xlsx_export_has_expected_safe_sheets():
    insight = make_insight()

    content = build_xlsx_export(
        run_id=uuid4(),
        dataset_version_id=uuid4(),
        status="completed",
        created_at="2026-10-06T12:00:00",
        ranked_insights=[
            insight
        ],
        explained_insights=[
            make_explained(
                insight
            )
        ],
        recommended_insights=[
            make_recommended(
                insight
            )
        ],
    )

    workbook = load_workbook(
        io.BytesIO(content)
    )

    assert workbook.sheetnames == [
        "Executive Summary",
        "Verified Insights",
        "Evidence",
        "Run Metadata",
    ]

    metadata = workbook[
        "Run Metadata"
    ]

    values = {
        row[0].value:
        row[1].value
        for row in metadata.iter_rows(
            min_row=2
        )
    }

    assert (
        values[
            "Raw Dataset Included"
        ]
        == "No"
    )

    assert (
        values[
            "Protected / PII Columns Included"
        ]
        == "No"
    )


def test_export_builder_does_not_need_raw_analysis_results():
    run = SimpleNamespace(
        ranked_insights=[
            make_insight()
        ],
        explained_insights=[],
        recommended_insights=[],
        analysis_results=[
            {
                "raw_customer_email":
                    "secret@example.com"
            }
        ],
    )

    content = build_csv_export(
        run.ranked_insights,
        run.explained_insights,
        run.recommended_insights,
    )

    assert (
        b"secret@example.com"
        not in content
    )



def test_pdf_export_is_valid_pdf_and_uses_verified_reporting_layer():
    insight = make_insight()

    content = build_pdf_export(
        run_id=uuid4(),
        dataset_version_id=uuid4(),
        status="completed",
        created_at=(
            "2026-10-06T12:00:00"
        ),
        ranked_insights=[
            insight
        ],
        explained_insights=[
            make_explained(
                insight
            )
        ],
        recommended_insights=[
            make_recommended(
                insight
            )
        ],
    )

    assert content.startswith(
        b"%PDF"
    )

    assert len(content) > 2000


def test_pdf_export_supports_empty_verified_results():
    content = build_pdf_export(
        run_id=uuid4(),
        dataset_version_id=uuid4(),
        status="completed",
        created_at=(
            "2026-10-06T12:00:00"
        ),
        ranked_insights=[],
        explained_insights=[],
        recommended_insights=[],
    )

    assert content.startswith(
        b"%PDF"
    )

    assert len(content) > 1000
